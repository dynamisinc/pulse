#!/usr/bin/env python3
"""
build-kestrel-msel.py: builds docs/demo/pack/kestrel-msel.xlsx, the Kestrel County boil-water run sheet rendered
as a Cadence MSEL, and then checks the file it wrote against Pulse's MSEL importer rules.

WHAT THE FILE IS FOR (demo-feedback DF-16, task C2 of #511)
    On stage the controller imports this workbook into the Pulse run sheet (Run sheet > Import, dry run first).
    Pulse keeps only the rows whose Delivery Method is social media (DF-9) and turns each into a single-post
    run-sheet item; every phone call, email, radio call, in-person (Verbal), written and EAS / WEA row is reported
    as "skipped", which is the point of the demo beat: Pulse reads the same MSEL the exercise already has.

WHAT IT CONTAINS
    * One StartEx row (the earliest Scenario Day / Scenario Time in the file) and an EndEx row.
    * One "Social Media" row for each of the 8 single-post run-sheet items that can travel through an MSEL:
      rs01, rs05, rs07, rs08, rs09, rs12, rs13 and rs15. Their persona, post text, media, notes and planned minute
      are READ FROM pack.json, so the MSEL cannot drift from the pack. Their Title, To / Target, Track, Expected
      Action, Priority and Location are MSEL-only and live in SOCIAL_ROWS below.
      rs03 (the impersonator post that the rs04 burst replies to) and rs10 (the verified utility's reply inside
      rs03's thread) are NOT in the MSEL: an MSEL row has no reply-to, and rs03 must exist before rs04 and rs10 can
      reply to it. They stay seeded, like every burst.
    * About 25 non-social rows (OTHER_ROWS below) around them: utility crew radio traffic, police, hospital, 311,
      school district, the state drinking-water program, a WEA offer and a contingency WEA, lab results, mutual aid.

TIMING (DF-11)
    The importer sets each item's PlannedMinute to the minutes from the earliest (Scenario Day, Scenario Time) of any
    data row (the StartEx row) to the row's own. Every social row is placed at StartEx + the item's plannedMinute in
    pack.json, so the import reproduces the authored T+N. Scheduled Time (Cadence's wall clock) is set to the same
    clock time: real-time play, which is also what the importer falls back to if Scenario Time is ever removed.

INJECT NUMBERS
    Sequential integers in time order across all rows (ties: the non-social row first). Each of the 8 pack items
    carries "msel": "<its Inject #>" in pack.json (the seed's -RunSheet exceptMsel leaves those items to the import).
    This script FAILS if pack.json's msel values do not match the numbers it assigns, and prints the right ones, so
    re-timing a row means updating pack.json in the same change.

CADENCE LAYOUT (MselWorksheetBuilder / ExcelFormattingHelper.MselColumns, read 2026-10-15)
    Sheet "MSEL"; columns, in order: Inject #, Title, Description, Scheduled Time, Scenario Day, Scenario Time,
    From / Source, To / Target, Delivery Method, Track, Phase, Expected Action, Notes, Priority, Location,
    Responsible Controller; then the optional Pulse extension column "Pulse Media". Inject # is a NUMBER cell;
    everything else is text, with times as "HH:mm" and Scenario Day / Priority as digit strings, exactly as Cadence
    writes them. Header bold on light blue, centred, wrapped, Cadence's column widths; alternate rows AliceBlue;
    autofilter. Added here (not in Cadence): a frozen header row. Cadence's "Phases" and "Objectives" sheets are
    included as an export with those options would have them.
    Cadence's delivery methods are Verbal, Phone, Email, Radio, Written, Simulation and Other; "Social Media" and
    "EAS / WEA" are Other with a custom name, which Cadence exports as that name.

RESPONSIBLE CONTROLLER
    Blank on the social rows, on purpose. The importer matches it against the exercise's staff display names: a
    blank cell is silent (the item is created unassigned), while any name that is not a staff display name on the
    target exercise adds a warning to that row. The non-social rows carry SimCell role names because the importer
    skips them before it reads that column.

PULSE MEDIA
    The media-library file name: the base name of the pack file, which is what Seed-DemoContent.ps1 uploads and
    what the server keeps (MediaUploadService.SanitizeFileName). rs01 news-mill-road.mp4, rs05 scoop-brown-tap.mp4,
    rs08 pio-statement-card.png. The seed must have uploaded the media before the import, or those rows fail.

THE CHECK (runs after every build, or alone with --check)
    Re-opens the workbook with openpyxl and applies a Python mirror of the importer (MselWorkbookReader +
    MselImportPlanner, src/Pulse.WebApi/Features/Injects/Import/): sheet choice, required headers, the cell-to-text
    rules, the social filter, DF-11 minutes, persona by handle, text length in code points after trimming (and no
    "<", so PostSanitizer cannot change a word), Title <= 120, Notes <= 500, unique Inject #, Pulse Media names
    against the pack's library files (<= 4 images or exactly 1 video), Responsible Controller warnings. It also
    checks that each social row's text, persona, media and minute equal the pack item, and prints the dry-run
    report the console would show: 34 rows, 8 social, 8 created, 26 skipped, 0 failed, 0 warnings on an exercise
    that has none of these Inject # values yet.

RUN
    python3 docs/demo/pack/tools/build-kestrel-msel.py            # build, then check
    python3 docs/demo/pack/tools/build-kestrel-msel.py --check    # check the committed file only
    Needs Python 3.10+ and openpyxl (pip install openpyxl). The output is byte-stable: re-running with unchanged
    inputs writes an identical file (fixed document timestamps and zip entry times).
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

PACK_DIR = Path(__file__).resolve().parent.parent
PACK_JSON = PACK_DIR / "pack.json"
OUTPUT = PACK_DIR / "kestrel-msel.xlsx"

# ---- the MSEL's clock -------------------------------------------------------------------------------------------
SCENARIO_DAY = 1
START_HOUR, START_MINUTE = 13, 0  # StartEx, scenario and scheduled (real-time play)
FIXED_STAMP = datetime(2026, 10, 15, 12, 0, 0)  # document created/modified, for a byte-stable file

# ---- Cadence's MSEL columns, in order, with its widths, then the Pulse extension ---------------------------------
CADENCE_COLUMNS = [
    ("Inject #", 10), ("Title", 40), ("Description", 60), ("Scheduled Time", 15), ("Scenario Day", 12),
    ("Scenario Time", 15), ("From / Source", 20), ("To / Target", 20), ("Delivery Method", 18), ("Track", 15),
    ("Phase", 20), ("Expected Action", 50), ("Notes", 40), ("Priority", 10), ("Location", 20),
    ("Responsible Controller", 20),
]
PULSE_MEDIA_COLUMN = ("Pulse Media", 24)
COLUMNS = CADENCE_COLUMNS + [PULSE_MEDIA_COLUMN]

LIGHT_BLUE = "FFADD8E6"  # ClosedXML XLColor.LightBlue (MSEL header)
ALICE_BLUE = "FFF0F8FF"  # XLColor.AliceBlue (even rows)
LIGHT_GREEN = "FF90EE90"  # XLColor.LightGreen (Phases header)
LIGHT_CORAL = "FFF08080"  # XLColor.LightCoral (Objectives header)

SOCIAL_MEDIA = "Social Media"
PULSE_LOCATION = "Pulse"
PUBLIC_TARGET = "Public / PIO"

# ---- phases (minutes from StartEx) and objectives ----------------------------------------------------------------
PHASES = [
    (1, "Second Main Break", "A second main breaks on Mill Road; the utility and the EOC assess it.", 0, 5),
    (2, "Misinformation", "A lookalike account claims boiling is unsafe and the advisory covers all of Fairhaven.", 6, 9),
    (3, "Public Information Response", "The county and the utility correct the record on verified accounts.", 10, 18),
    (4, "Sustained Operations", "Restoration, sampling and distribution continue.", 19, 45),
]
OBJECTIVES = [
    (1, "Public Information and Warning",
     "Issue accurate, consistent public messages on the boil water advisory and correct misinformation within "
     "15 minutes of detection."),
    (2, "Operational Coordination",
     "Coordinate Kestrel County, Fairhaven Water and partners through the EOC and the Joint Information Center."),
    (3, "Critical Infrastructure: Water",
     "Track both main breaks, sampling and restoration, and share their status with the EOC and the public."),
    (4, "Mass Care Services",
     "Keep bottled water distribution running, including delivery to residents with access and functional needs."),
]

# ---- the 8 social rows: MSEL-only fields. Persona, text, media, notes and minute come from pack.json ---------------
SOCIAL_ROWS = {
    "rs01": dict(
        title="Metro Desk 9 breaking video: second main break on Mill Road",
        expected="JIC confirms the second break with the utility before amplifying; PIO prepares a holding statement.",
        priority="2"),
    "rs05": dict(
        title="The Scoop amplifies the rumor with a brown-tap video",
        expected="JIC tracks the rumor's spread and sources; PIO prepares the correction.",
        priority="2"),
    "rs07": dict(
        title="Media query: Metro Desk 9 asks officials to confirm (cue the PIO)",
        expected="PIO posts the official correction as @KestrelCoEM and attaches the notice card.",
        priority="1"),
    "rs08": dict(
        title="BACKUP: County EM official statement with the notice card",
        expected="None. Controller backup: fire only if the PIO's live post fails, otherwise skip.",
        priority="3"),
    "rs09": dict(
        title="Resident asks County EM about formula and ice",
        expected="PIO answers with Public Health's guidance (formula made with boiled water is fine; throw out ice).",
        priority="3"),
    "rs12": dict(
        title="Metro Desk 9 reports the official correction",
        expected="JIC confirms the media carried the correction accurately.",
        priority="3"),
    "rs13": dict(
        title="OPTIONAL: impersonator doubles down (second takedown target)",
        expected="Controller takes the post down from Live world if showing takedown again; JIC reports the account.",
        priority="4"),
    "rs15": dict(
        title="Fairhaven Water: Mill Road break isolated (then Freeze)",
        expected="PIO shares the utility update. Controller freezes the world after it posts.",
        priority="3"),
}
SOCIAL_TRACK = "Public Information"

# ---- the non-social rows ------------------------------------------------------------------------------------------
# minute = minutes from StartEx. "{inject:<id>}" in a description or note becomes that row's Inject #.
OTHER_ROWS = [
    dict(id="startex", minute=0, method="Verbal", title="StartEx",
         description=(
             "Exercise Director announces StartEx. Situation: an overnight water main break on Elm Street near the "
             "Eastside pump station; Fairhaven Water's precautionary boil water advisory is in effect for Zones 2, 3 "
             "and 4; free bottled water at Elm Street Fire Station and Fairhaven High School; the Kestrel County EOC "
             "is partially activated (Level 2) with a Joint Information Center."),
         source="Exercise Director", target="All players", track="Exercise Control",
         expected="Players take their positions and begin play from the situation brief.",
         notes="The Pulse opening feed is already seeded: about six scenario hours of posts before StartEx.",
         priority="3", location="Kestrel County EOC", controller="Lead Controller"),
    dict(id="crew-leak", minute=1, method="Radio", title="Utility crew reports a second leak on Mill Road",
         description=(
             "Fairhaven Water Crew 3 to Water Operations: water is coming up through the pavement on Mill Road near "
             "the Eastside reservoir and running to the storm drains. Heavy flow. Requesting a second crew and "
             "authorization to close the Mill Road feed valve."),
         source="Fairhaven Water Crew 3", target="Fairhaven Water Operations", track="Utility",
         expected="Utility operations dispatches a second crew, informs the EOC utility liaison and assesses the "
                  "effect on Zone 4 pressure.",
         notes="SimCell plays the crew on the utility radio net.",
         priority="1", location="Mill Road", controller="SimCell - Utility"),
    dict(id="pd-closure", minute=2, method="Phone", title="Police close Mill Road",
         description=(
             "Fairhaven Police dispatch calls the EOC: officers have closed Mill Road between Birch Lane and "
             "Reservoir Drive because of water over the roadway, and traffic is detoured to Cedar Avenue. Asks "
             "whether the utility expects the closure to last past the evening commute."),
         source="Fairhaven Police Dispatch", target="EOC Operations Section Chief", track="Law Enforcement",
         expected="EOC logs the closure, shares it with the JIC and confirms the expected duration with the utility.",
         notes="", priority="2", location="Kestrel County EOC", controller="SimCell - EOC"),
    dict(id="utility-confirms", minute=4, method="Phone", title="Utility confirms the second break; advisory area unchanged",
         description=(
             "Fairhaven Water's operations manager to the EOC Director: the Mill Road leak is a second main break. "
             "Zone 4 pressure is dropping. The boil water advisory area does not change for now: Zones 2, 3 and 4 "
             "only. The utility asks that every public message says so plainly."),
         source="Fairhaven Water Operations Manager", target="EOC Director", track="Utility",
         expected="EOC Director confirms the advisory boundary with the PIO; the JIC prepares a holding statement on "
                  "the second break.",
         notes="", priority="1", location="Kestrel County EOC", controller="SimCell - Utility"),
    dict(id="state-dw", minute=5, method="Email", title="State drinking water program: sampling requirements",
         description=(
             "The State Drinking Water Program acknowledges Fairhaven Water's boil water advisory and reminds the "
             "utility that it stays in effect until two consecutive rounds of bacteriological samples from the "
             "affected zones are clean. Requests the sampling plan for the Mill Road segment and the time of the next "
             "sample round."),
         source="State Drinking Water Program", target="Fairhaven Water; EOC Utility Liaison", track="Utility",
         expected="Utility sends the sampling plan; the EOC gives the PIO the conditions for lifting the advisory so "
                  "public messages stay consistent.",
         notes="", priority="3", location="Fairhaven Water Operations Center", controller="SimCell - State"),
    dict(id="hospital", minute=6, method="Phone", title="Hospital asks whether it is inside the advisory",
         description=(
             "Kestrel Regional Medical Center infection prevention calls the EOC: the hospital sits on the edge of "
             "Zone 2. Is it inside the advisory, is the dialysis unit's water treatment affected, and who will tell "
             "the hospital when the advisory changes?"),
         source="Kestrel Regional Medical Center", target="EOC Public Health Liaison", track="Public Health",
         expected="Public Health confirms the hospital's zone and notification path and asks the utility to add it "
                  "to critical-customer outreach.",
         notes="", priority="2", location="Kestrel County EOC", controller="SimCell - Health"),
    dict(id="311", minute=7, method="Phone", title="311 call volume spikes: is boiling unsafe?",
         description=(
             "The Kestrel County 311 center reports a sudden spike in calls. Residents are asking whether boiling tap "
             "water is now unsafe and whether the advisory covers all of Fairhaven. Call takers have no approved "
             "answer and ask the JIC for talking points."),
         source="Kestrel County 311 Supervisor", target="JIC / PIO", track="Public Information",
         expected="JIC gives 311 approved talking points (boiling is safe; Zones 2-4 only) and starts tracing where "
                  "the claim comes from.",
         notes="Deliver after the lookalike's \"do NOT boil\" post has appeared in Pulse.",
         priority="1", location="Joint Information Center", controller="SimCell - EOC"),
    dict(id="site-elm", minute=9, method="Radio", title="Distribution site: residents asking if boiling is unsafe",
         description=(
             "Elm Street Fire Station site lead: the line has doubled in ten minutes. Residents are showing "
             "volunteers a post that says boiling is unsafe and asking whether to stop. Volunteers need one sentence "
             "to tell them."),
         source="Elm Street Fire Station Site Lead", target="EOC Logistics; JIC", track="Mass Care",
         expected="JIC sends site volunteers the same approved line; Logistics checks site stock against the surge.",
         notes="", priority="2", location="Elm Street Fire Station", controller="SimCell - Mass Care"),
    dict(id="health-guidance", minute=11, method="Email", title="Public Health guidance: formula, ice and showering",
         description=(
             "Kestrel County Public Health sends the PIO its standing boil water guidance: use boiled (then cooled) "
             "or bottled water for drinking, cooking, brushing teeth and infant formula; formula made earlier with "
             "boiled water is fine; throw out ice made from tap water during the advisory; showering is fine if no "
             "water is swallowed, and a sponge bath is safer for infants."),
         source="Kestrel County Public Health", target="PIO / JIC", track="Public Health",
         expected="PIO uses the guidance to answer residents' follow-up questions.",
         notes="Gives the PIO the answer to the resident's formula-and-ice question that follows.",
         priority="2", location="Joint Information Center", controller="SimCell - Health"),
    dict(id="wea-offer", minute=13, method="Phone", title="State Warning Point offers a WEA",
         description=(
             "The State Warning Point calls the EOC Director: with the boiling rumor spreading, does the county want "
             "to send a Wireless Emergency Alert (Public Safety Message) to Zones 2-4? The warning point needs the "
             "approved text and the alert area within 15 minutes."),
         source="State Warning Point", target="EOC Director; PIO", track="Public Information",
         expected="EOC Director decides; if yes, the PIO drafts the WEA text for Zones 2-4.",
         notes="Contingency: Inject {inject:wea-release} delivers the alert only if players ask for one here.",
         priority="2", location="Kestrel County EOC", controller="SimCell - State"),
    dict(id="mayor", minute=14, method="Phone", title="Mayor's office asks for a briefing",
         description=(
             "The Mayor's chief of staff calls: council members are forwarding posts that say the advisory now covers "
             "the whole city and that boiling is unsafe. The Mayor wants a two-minute briefing and the county's "
             "statement before speaking to reporters."),
         source="Office of the Mayor, City of Fairhaven", target="EOC Director", track="EOC Management",
         expected="EOC Director briefs the Mayor's office and shares the official statement; the JIC aligns the "
                  "Mayor's talking points.",
         notes="", priority="2", location="Kestrel County EOC", controller="SimCell - EOC"),
    dict(id="crew-isolated", minute=17, method="Radio", title="Utility crew: Mill Road feed isolated",
         description=(
             "Fairhaven Water Crew 3: the Mill Road feed valve is closed and the second break is isolated. Starting "
             "excavation. Zone 4 will see low pressure for a few more hours."),
         source="Fairhaven Water Crew 3", target="Fairhaven Water Operations", track="Utility",
         expected="Utility updates the EOC and posts a public update.",
         notes="Sets up the utility's closing post (the next Social Media row).",
         priority="3", location="Mill Road", controller="SimCell - Utility"),
    dict(id="schools", minute=19, method="Email", title="School district asks about classes tomorrow",
         description=(
             "Fairhaven Unified School District asks whether its four schools in Zones 2-4 should open tomorrow, and "
             "whether drinking fountains and cafeteria kitchens can be used. The district wants a decision in time "
             "for its message to families tonight."),
         source="Fairhaven Unified School District", target="EOC Public Health Liaison", track="Schools",
         expected="Public Health advises (fountains off, bottled water, kitchen guidance); the JIC coordinates the "
                  "district's message to families.",
         notes="", priority="3", location="Kestrel County EOC", controller="SimCell - Health"),
    dict(id="senior-center", minute=20, method="Phone", title="Senior center requests bottled water delivery",
         description=(
             "The Eastside Senior Center director calls: about 40 residents of the apartments next door cannot get to "
             "a distribution site. Can the county deliver bottled water today?"),
         source="Eastside Senior Center", target="EOC Mass Care (ESF-6)", track="Mass Care",
         expected="Mass Care arranges a delivery and adds the building to the access and functional needs list.",
         notes="", priority="2", location="Kestrel County EOC", controller="SimCell - Mass Care"),
    dict(id="lab-round-1", minute=21, method="Email", title="Lab results: first sample round, Zone 2",
         description=(
             "Fairhaven Water Quality Laboratory: round 1 bacteriological samples from Zone 2 are complete; all 12 "
             "are absent for total coliform. Round 2 will be collected tomorrow morning. Zones 3 and 4 will be sampled "
             "once pressure is restored. The advisory stays in effect until both rounds are clean."),
         source="Fairhaven Water Quality Laboratory", target="Fairhaven Water; EOC Utility Liaison", track="Utility",
         expected="Utility and EOC note the result. The PIO does not announce an end to the advisory: one clean round "
                  "is not enough.",
         notes="", priority="3", location="Fairhaven Water Operations Center", controller="SimCell - Utility"),
    dict(id="mutual-aid", minute=22, method="Email", title="Neighboring water district offers mutual aid",
         description=(
             "Kestrel Valley Water District offers a crew, a valve truck and a 4,000-gallon potable water tanker "
             "under the county mutual aid agreement. Asks where to stage and who to report to."),
         source="Kestrel Valley Water District", target="EOC Logistics", track="Utility",
         expected="Logistics accepts or declines, names a staging area and a point of contact, and logs the resource.",
         notes="", priority="3", location="Kestrel County EOC", controller="SimCell - Utility"),
    dict(id="911", minute=23, method="Phone", title="911 center: discolored-water calls from Zone 4",
         description=(
             "Kestrel County Emergency Communications reports a run of non-emergency 911 calls from Zone 4 about "
             "brown water since the Mill Road break. Asks whether to keep routing them to the utility's customer line."),
         source="Kestrel County Emergency Communications", target="EOC Operations Section Chief",
         track="Emergency Communications",
         expected="EOC confirms the routing and sends the 911 center the discolored-water talking points.",
         notes="", priority="4", location="Kestrel County EOC", controller="SimCell - EOC"),
    dict(id="dialysis", minute=24, method="Phone", title="Dialysis clinic asks about Zone 4 pressure",
         description=(
             "The East Side Dialysis Center calls: its treatment system handles the advisory, but low pressure in "
             "Zone 4 could interrupt afternoon treatments. Asks for a pressure estimate and priority restoration."),
         source="East Side Dialysis Center", target="Fairhaven Water; EOC Public Health Liaison", track="Public Health",
         expected="Utility logs the clinic as a critical customer; Public Health checks whether patients need to move "
                  "to another clinic.",
         notes="", priority="2", location="Fairhaven Water Operations Center", controller="SimCell - Health"),
    dict(id="wea-release", minute=25, method="EAS / WEA", title="WEA reaches Zones 2-4 (contingency)",
         description=(
             "CONTINGENCY, only if players requested a WEA at Inject {inject:wea-offer}: a Wireless Emergency Alert "
             "reaches phones in Zones 2-4 with the players' approved text. If they approved the alert but wrote no "
             "text, use: \"Kestrel Co EM: Fairhaven Zones 2-4 boil water advisory continues. Boiling is safe.\""),
         source="State Warning Point", target="Residents of Zones 2-4", track="Public Information",
         expected="PIO confirms the alert went out as approved and repeats the message on the official accounts.",
         notes="Skip if no WEA was requested.",
         priority="2", location="Zones 2-4", controller="SimCell - State"),
    dict(id="site-high-school", minute=26, method="Radio", title="High school site needs resupply",
         description=(
             "Fairhaven High School site lead: about 300 cases left at the current pace, roughly an hour. Requests a "
             "resupply and two more volunteers for the infant and accessibility table."),
         source="Fairhaven High School Site Lead", target="EOC Logistics", track="Mass Care",
         expected="Logistics schedules the resupply and volunteers; the JIC keeps site information accurate.",
         notes="", priority="2", location="Fairhaven High School", controller="SimCell - Mass Care"),
    dict(id="sitrep", minute=28, method="Written", title="Planning Section requests SITREP inputs",
         description=(
             "EOC Planning Section memo: section and ESF inputs for the next situation report are due before the "
             "operational period briefing. Include the second break, sampling results, distribution counts, the "
             "rumor and the county's response."),
         source="EOC Planning Section Chief", target="All EOC sections; JIC", track="EOC Management",
         expected="Sections submit inputs; the JIC includes the misinformation timeline and the actions taken.",
         notes="", priority="4", location="Kestrel County EOC", controller="Lead Controller"),
    dict(id="chamber", minute=30, method="Phone", title="Chamber of Commerce asks about restaurant guidance",
         description=(
             "Fairhaven Chamber of Commerce calls: east side restaurants and coffee shops want to know whether they "
             "can stay open, and what to do about ice machines, soda fountains and handwashing."),
         source="Fairhaven Chamber of Commerce", target="EOC Public Health Liaison", track="Public Health",
         expected="Public Health provides food-service guidance; the JIC posts a summary for businesses.",
         notes="", priority="4", location="Kestrel County EOC", controller="SimCell - Health"),
    dict(id="utility-gm", minute=33, method="Verbal", title="Utility general manager asks for a joint briefing",
         description=(
             "Fairhaven Water's general manager arrives at the EOC and asks to plan a joint media briefing with the "
             "county within the hour: who speaks, what is said about the two breaks, and how the advisory will end."),
         source="Fairhaven Water General Manager", target="EOC Director; PIO", track="Public Information",
         expected="EOC Director and PIO agree speakers, key messages and timing for a joint briefing.",
         notes="Delivered in person by the controller playing the general manager.",
         priority="3", location="Kestrel County EOC", controller="Lead Controller"),
    dict(id="sheriff", minute=36, method="Phone", title="Sheriff: free water being resold near Elm Street",
         description=(
             "Kestrel County Sheriff's Office reports complaints that people near Elm Street are reselling the free "
             "bottled water for cash. Asks whether the county wants deputies at the sites and a public reminder that "
             "the water is free."),
         source="Kestrel County Sheriff's Office", target="EOC Law Enforcement Liaison", track="Law Enforcement",
         expected="EOC decides on a deputy presence; the JIC repeats that the water is free and names the official "
                  "sites.",
         notes="", priority="3", location="Elm Street Fire Station", controller="SimCell - EOC"),
    dict(id="crew-repair", minute=40, method="Radio", title="Utility crew: repair clamp on, reopening Mill Road feed",
         description=(
             "Fairhaven Water Crew 3: repair clamp installed on the Mill Road main. Slowly reopening the feed. Expect "
             "discolored water in Zone 4 while the lines refill; hydrant flushing to follow."),
         source="Fairhaven Water Crew 3", target="Fairhaven Water Operations", track="Utility",
         expected="Utility plans flushing; the JIC prepares a discolored-water message for Zone 4.",
         notes="", priority="3", location="Mill Road", controller="SimCell - Utility"),
    dict(id="endex", minute=45, method="Verbal", title="EndEx",
         description="Exercise Director announces EndEx. Players stop play and stay in place for the hotwash.",
         source="Exercise Director", target="All players", track="Exercise Control",
         expected="Players stop play; controllers and evaluators collect notes for the hotwash.",
         notes="In the Pulse demo the world is frozen after the utility's closing post (the last Social Media row).",
         priority="3", location="Kestrel County EOC", controller="Lead Controller"),
]

# ---- importer limits (InjectItemValidator / MselImportPlanner / InjectItem) --------------------------------------
MAX_TEXT = 280
MAX_TITLE = 120
MAX_NOTES = 500
MAX_INJECT_NUMBER = 64
MAX_IMAGES = 4
REQUIRED_HEADERS = ["Inject #", "Delivery Method", "From / Source", "Description"]
KNOWN_HEADERS = [
    "Inject #", "Title", "Description", "Scheduled Time", "Scenario Day", "Scenario Time", "From / Source",
    "Delivery Method", "Notes", "Responsible Controller", "Pulse Media",
]


def clock(minute: int) -> str:
    """HH:mm for StartEx + minute (Cadence's TimeOnly.ToString("HH:mm"))."""
    total = START_HOUR * 60 + START_MINUTE + minute
    if total >= 24 * 60:
        raise SystemExit(f"minute {minute} runs past midnight; move StartEx or add Scenario Day handling")
    return f"{total // 60:02d}:{total % 60:02d}"


def phase_for(minute: int) -> str:
    for _, name, _, start, end in PHASES:
        if start <= minute <= end:
            return name
    raise SystemExit(f"no phase covers minute {minute}")


def load_pack() -> dict:
    with PACK_JSON.open(encoding="utf-8") as handle:
        return json.load(handle)


def social_rows_from_pack(pack: dict) -> list[dict]:
    """The 8 social rows, read from the pack's run sheet (persona, text, media, notes, minute)."""
    media_files = {entry["key"]: entry["file"] for entry in pack["media"]}
    items = {item["key"]: item for item in pack["runSheet"]["items"]}
    rows = []
    for key, meta in SOCIAL_ROWS.items():
        item = items.get(key)
        if item is None:
            raise SystemExit(f"pack.json has no run-sheet item {key}")
        if item["kind"] != "post" or len(item["posts"]) != 1:
            raise SystemExit(f"{key} must be a single-post item to travel through an MSEL")
        post = item["posts"][0]
        if "replyTo" in post:
            raise SystemExit(f"{key} is a reply; an MSEL row has no reply-to")
        media = [Path(media_files[ref["ref"]]).name for ref in post.get("media", [])]
        rows.append(dict(
            id=key, minute=item["plannedMinute"], method=SOCIAL_MEDIA, title=meta["title"],
            description=post["text"], source="@" + post["persona"], target=PUBLIC_TARGET, track=SOCIAL_TRACK,
            expected=meta["expected"], notes=item.get("notes", ""), priority=meta["priority"],
            location=PULSE_LOCATION, controller="", media=", ".join(media), pack_key=key))
    return rows


def ordered_rows(pack: dict) -> list[dict]:
    """All rows in time order (non-social first on a tie), numbered 1..N, with {inject:id} references filled in."""
    rows = [dict(row, media="", pack_key=None) for row in OTHER_ROWS] + social_rows_from_pack(pack)
    rows.sort(key=lambda row: (row["minute"], row["method"] == SOCIAL_MEDIA))
    numbers = {row["id"]: index for index, row in enumerate(rows, start=1)}
    if len(numbers) != len(rows):
        raise SystemExit("row ids must be unique")

    def fill(text: str) -> str:
        return re.sub(r"\{inject:([^}]+)\}", lambda m: str(numbers[m.group(1)]), text)

    for row in rows:
        row["number"] = numbers[row["id"]]
        row["description"] = fill(row["description"])
        row["notes"] = fill(row["notes"])
    if rows[0]["id"] != "startex" or rows[0]["minute"] != 0:
        raise SystemExit("the StartEx row must be the earliest row")
    return rows


def verify_pack_msel(pack: dict, rows: list[dict]) -> None:
    """pack.json's "msel" on each item must equal the Inject # this script gives its row, and only those 8 items."""
    expected = {row["pack_key"]: str(row["number"]) for row in rows if row["pack_key"]}
    problems = []
    for item in pack["runSheet"]["items"]:
        have = item.get("msel")
        want = expected.get(item["key"])
        if have != want:
            problems.append(f'  {item["key"]}: pack has {have!r}, the MSEL gives {want!r}')
    if problems:
        print("pack.json's run-sheet \"msel\" values do not match the MSEL's Inject # values:", file=sys.stderr)
        print("\n".join(problems), file=sys.stderr)
        raise SystemExit(1)


# ---- writing ------------------------------------------------------------------------------------------------------
def header_row(sheet, headers_and_widths, fill_hex: str, wrap: bool) -> None:
    for column, (header, width) in enumerate(headers_and_widths, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=fill_hex)
        cell.alignment = Alignment(horizontal="center", wrap_text=wrap)
        sheet.column_dimensions[cell.column_letter].width = width


def build_workbook(rows: list[dict]) -> bytes:
    workbook = Workbook()
    msel = workbook.active
    msel.title = "MSEL"
    header_row(msel, COLUMNS, LIGHT_BLUE, wrap=True)
    for offset, row in enumerate(rows):
        sheet_row = offset + 2
        values = [
            row["number"],  # a NUMBER cell, as Cadence writes InjectNumber
            row["title"], row["description"], clock(row["minute"]), str(SCENARIO_DAY), clock(row["minute"]),
            row["source"], row["target"], row["method"], row["track"], phase_for(row["minute"]), row["expected"],
            row["notes"], row["priority"], row["location"], row["controller"], row["media"],
        ]
        for column, value in enumerate(values, start=1):
            cell = msel.cell(row=sheet_row, column=column, value=value)
            if column > 1:
                cell.data_type = "s"  # text, never a formula or a number (Cadence writes strings)
            if sheet_row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ALICE_BLUE)
    msel.auto_filter.ref = msel.dimensions
    msel.freeze_panes = "A2"

    phases = workbook.create_sheet("Phases")
    header_row(phases, [("Sequence", 10), ("Name", 30), ("Description", 50), ("Start Time", 12), ("End Time", 12)],
               LIGHT_GREEN, wrap=False)
    for offset, (sequence, name, description, start, end) in enumerate(PHASES, start=2):
        for column, value in enumerate([sequence, name, description, clock(start), clock(end)], start=1):
            phases.cell(row=offset, column=column, value=value)
    phases.auto_filter.ref = phases.dimensions

    objectives = workbook.create_sheet("Objectives")
    header_row(objectives, [("Objective #", 12), ("Name", 40), ("Description", 60)], LIGHT_CORAL, wrap=False)
    for offset, (number, name, description) in enumerate(OBJECTIVES, start=2):
        for column, value in enumerate([number, name, description], start=1):
            objectives.cell(row=offset, column=column, value=value)
    objectives.auto_filter.ref = objectives.dimensions

    workbook.properties.creator = "Pulse demo pack (build-kestrel-msel.py)"
    workbook.properties.title = "Kestrel County boil water MSEL"
    workbook.properties.created = FIXED_STAMP
    buffer = io.BytesIO()
    workbook.save(buffer)
    return stabilize(buffer.getvalue())


def stabilize(xlsx: bytes) -> bytes:
    """Fix the save-time stamp openpyxl writes and every zip entry time, so the same inputs give the same bytes."""
    stamp = FIXED_STAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    source = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              lambda m: m.group(1) + stamp.encode() + m.group(2), data)
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_STAMP.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, data)
    return out.getvalue()


# ---- checking: a Python mirror of MselWorkbookReader + MselImportPlanner -------------------------------------------
TIME_24 = re.compile(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$")
TIME_12 = re.compile(r"^(\d{1,2})(?::(\d{2}))?(?::(\d{2}))?\s?([AaPp][Mm])$")


def cell_text(value) -> str:
    """MselWorkbookReader.CellText: blank -> "", number -> invariant shortest form, text as is."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(int(value)) if float(value).is_integer() else repr(float(value))
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def parse_time_of_day(cell: str):
    """MselImportPlanner.ParseTimeOfDay for the forms this file can hold (H:mm, HH:mm, :ss, h:mm tt)."""
    value = cell.strip()
    if not value:
        return None
    match = TIME_24.match(value)
    if match:
        hour, minute = int(match.group(1)), int(match.group(2))
        return hour * 60 + minute if hour < 24 and minute < 60 else None
    match = TIME_12.match(value)
    if match:
        hour, minute = int(match.group(1)), int(match.group(2) or 0)
        if not 1 <= hour <= 12 or minute >= 60:
            return None
        hour = hour % 12 + (12 if match.group(4).lower() == "pm" else 0)
        return hour * 60 + minute
    return None


def parse_day(cell: str):
    value = cell.strip()
    if not value:
        return 1, None
    try:
        number = float(value)
    except ValueError:
        return 1, f"Scenario Day '{value}' is not a whole number; day 1 assumed."
    if number.is_integer() and 0 <= number <= 10_000:
        return int(number), None
    return 1, f"Scenario Day '{value}' is not a whole number; day 1 assumed."


def is_social(method: str) -> bool:
    value = method.strip()
    return "social" in value.lower() or value.lower() == "pulse"


def clean(value: str):
    """MselImportPlanner.Clean. The check refuses any '<' up front, so PostSanitizer would change nothing."""
    value = value.strip()
    return value or None


def check(path: Path, pack: dict, expected_rows: list[dict] | None) -> int:
    workbook = load_workbook(path)  # values, never formulas (the file has none)
    sheet = next((ws for ws in workbook.worksheets if ws.title.strip().lower() == "msel"), workbook.worksheets[0])
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    rows = [(index, [cell_text(value) for value in row]) for index, row in enumerate(rows, start=1)]
    rows = [(index, cells) for index, cells in rows if any(cell.strip() for cell in cells)]
    header_index, headers = rows[0]
    data = rows[1:]

    problems: list[str] = []
    if [header for header, _ in COLUMNS] != headers:
        problems.append(f"header row is not Cadence's layout + Pulse Media: {headers}")
    columns: dict[str, int] = {}
    for position, header in enumerate(headers):
        known = next((name for name in KNOWN_HEADERS if name.lower() == header.strip().lower()), None)
        if known and known not in columns:
            columns[known] = position
    missing = [header for header in REQUIRED_HEADERS if header not in columns]
    if missing:
        problems.append(f"missing required column(s): {missing}")

    def get(cells, header):
        position = columns.get(header)
        return cells[position] if position is not None and position < len(cells) else ""

    # DF-11: Scenario Time if any row has one, else Scheduled Time.
    use_scenario = any(parse_time_of_day(get(cells, "Scenario Time")) is not None for _, cells in data)
    time_header = "Scenario Time" if use_scenario else "Scheduled Time"
    keys = {}
    for index, cells in data:
        raw = get(cells, time_header)
        minutes = parse_time_of_day(raw)
        day, day_warning = parse_day(get(cells, "Scenario Day"))
        if minutes is None:
            problems.append(f"row {index}: {time_header} {raw!r} does not parse; the row would be unplanned")
        if day_warning:
            problems.append(f"row {index}: {day_warning}")
        keys[index] = None if minutes is None else day * 24 * 60 + minutes
    start = min(key for key in keys.values() if key is not None)

    handles = {persona["handle"].lower(): persona["handle"] for persona in pack["personas"]}
    library = {}
    for entry in pack["media"]:  # what the seed uploads; poster frames are excluded from the importer's list
        library[Path(entry["file"]).name.lower()] = entry["kind"]
    pack_items = {item["key"]: item for item in pack["runSheet"]["items"]}
    by_number = {item["msel"]: item for item in pack["runSheet"]["items"] if "msel" in item}

    seen: dict[str, int] = {}
    report = []
    counts = {"social": 0, "created": 0, "skipped": 0, "failed": 0, "warnings": 0}
    for index, cells in data:
        number = clean(get(cells, "Inject #"))
        minute = None if keys[index] is None else keys[index] - start
        method = get(cells, "Delivery Method")
        if any("<" in cell for cell in cells):
            problems.append(f"row {index}: a cell contains '<' (the sanitizer could change it)")
        if not is_social(method):
            counts["skipped"] += 1
            report.append(f"  row {index:>2}  #{number:<3} skipped   T+{minute:<3} {method}")
            continue
        counts["social"] += 1
        failure = None
        warnings = []
        source = clean(get(cells, "From / Source")) or ""
        handle = source.lstrip("@").strip()
        text = clean(get(cells, "Description"))
        title = clean(get(cells, "Title")) or f"Inject {number}"
        notes = clean(get(cells, "Notes"))
        media_cell = clean(get(cells, "Pulse Media"))
        controller = clean(get(cells, "Responsible Controller"))
        if number is None:
            failure = "Inject # is empty"
        elif len(number) > MAX_INJECT_NUMBER:
            failure = "Inject # too long"
        elif number.lower() in seen:
            failure = f"duplicate of row {seen[number.lower()]}"
        else:
            seen[number.lower()] = index
        if not failure and handle.lower() not in handles:
            failure = f"unknown persona '{handle}'"
        if not failure and (text is None or len(text) > MAX_TEXT):
            failure = f"Description is {0 if text is None else len(text)} code points"
        if not failure and media_cell:
            kinds = []
            for name in [part.strip() for part in media_cell.split(",") if part.strip()]:
                if name.lower() not in library:
                    failure = f"unknown Pulse Media file '{name}'"
                    break
                kinds.append(library[name.lower()])
            images, videos = kinds.count("image"), kinds.count("video")
            if not failure and (images > MAX_IMAGES or videos > 1 or (videos and images)):
                failure = "too many media"
        if len(title) > MAX_TITLE:
            warnings.append("Title truncated")
        if notes and len(notes) > MAX_NOTES:
            warnings.append("Notes truncated")
        if controller:
            warnings.append(f"Responsible Controller '{controller}' must match a staff display name or it warns")
        if failure:
            counts["failed"] += 1
            problems.append(f"row {index}: would FAIL: {failure}")
        else:
            counts["created"] += 1
        counts["warnings"] += len(warnings)
        problems.extend(f"row {index}: warning: {warning}" for warning in warnings)

        # Equal to the pack item that carries this Inject # as its msel.
        item = by_number.get(number)
        if item is None:
            problems.append(f"row {index}: no pack run-sheet item has \"msel\": \"{number}\"")
        else:
            post = item["posts"][0]
            media = ", ".join(Path(next(m["file"] for m in pack["media"] if m["key"] == ref["ref"])).name
                              for ref in post.get("media", []))
            if text != post["text"]:
                problems.append(f"row {index}: Description differs from {item['key']}'s post text")
            if handle != post["persona"]:
                problems.append(f"row {index}: persona {handle} is not {item['key']}'s {post['persona']}")
            if (media_cell or "") != media:
                problems.append(f"row {index}: Pulse Media {media_cell!r} is not {item['key']}'s {media!r}")
            if minute != item["plannedMinute"]:
                problems.append(f"row {index}: T+{minute} is not {item['key']}'s plannedMinute {item['plannedMinute']}")
            if (notes or "") != item.get("notes", "").strip():
                problems.append(f"row {index}: Notes differ from {item['key']}'s notes")
        report.append(f"  row {index:>2}  #{number:<3} {'failed' if failure else 'created':<9} T+{minute:<3} "
                      f"@{handle} ({item['key'] if item else '?'}){' media: ' + media_cell if media_cell else ''}")

    if expected_rows is not None and len(expected_rows) != len(data):
        problems.append(f"{len(data)} data rows, expected {len(expected_rows)}")
    missing_items = [key for key, item in pack_items.items() if "msel" in item and item["msel"] not in seen]
    if missing_items:
        problems.append(f"pack items with an msel the workbook does not carry: {missing_items}")

    print(f"Dry-run mirror of {path.name} (sheet {sheet.title}, header row {header_index}, DF-11 on {time_header}, "
          f"StartEx = minute {start % (24 * 60)} of day {start // (24 * 60)}):")
    print("\n".join(report))
    print(f"  total {len(data)} rows: social {counts['social']}, created {counts['created']}, "
          f"skipped {counts['skipped']}, failed {counts['failed']}, warnings {counts['warnings']}")
    if problems:
        print("PROBLEMS:", file=sys.stderr)
        print("\n".join("  " + problem for problem in problems), file=sys.stderr)
        return 1
    print("OK: every row parses under the importer's rules and the social rows match pack.json.")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--check", action="store_true", help="only check the existing workbook")
    args = parser.parse_args()

    pack = load_pack()
    rows = ordered_rows(pack)
    verify_pack_msel(pack, rows)
    if not args.check:
        OUTPUT.write_bytes(build_workbook(rows))
        methods: dict[str, int] = {}
        for row in rows:
            methods[row["method"]] = methods.get(row["method"], 0) + 1
        print(f"Wrote {OUTPUT.relative_to(PACK_DIR.parent.parent.parent)}: {len(rows)} rows "
              f"({', '.join(f'{name} {count}' for name, count in sorted(methods.items()))})")
    return check(OUTPUT, pack, rows)


if __name__ == "__main__":
    sys.exit(main())
